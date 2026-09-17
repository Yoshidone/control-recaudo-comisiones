"""
Núcleo de negocio de "Control de Recaudo y Comisiones".

Este módulo es una copia FIEL de la lógica de cálculo del archivo original
app_pc.py (v31). No se modificó ninguna regla de negocio: mismos alias de
columnas, mismo parser streaming de XLSX, misma resolución de duplicados
PSP_TIN + Moneda, mismo formateo de Excel de salida.

Único cambio respecto al original: se retiró código que NUNCA se ejecutaba
en la app real (auditoría de eficiencia). La app de escritorio original
usaba exclusivamente el motor SQLite de más abajo (conn/calc_where/
main_where/query_*), y tenía además una segunda implementación en pandas
puro (process(), detail(), duplicate_rows(), accounting(), unique_sheet(),
export(), find_col(), EXCLUDE, dt()) que ningún botón de la interfaz
llamaba jamás. Se eliminó ese código muerto (~170 líneas, ~20% del
archivo) porque no aportaba nada salvo peso y confusión de mantenimiento;
el comportamiento de la app no cambia en absoluto.
"""

import os
import re
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
from xml.parsers import expat
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import load_workbook  # noqa: F401  (usado por quien importe este módulo para reabrir/estilar)
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TITLE = "Control de Recaudo y Comisiones"

ALIASES = {
    "com_id": ["com_id", "Com_ID", "com_public_id", "Com_Public_ID"],
    "comercio": ["Com_Nombre", "Com_Name", "com_nombre", "Comercio"],
    "fecha": ["TX_GMT_Peru", "FECHA", "Fecha"],
    "recaudo": ["PY_amount", "RECAUDO", "Recaudo"],
    "comision": ["SF_amount", "COMISION", "COMSIION", "Comision", "Comisión"],
    "metodo_pago": ["MétodoPago", "MetodoPago", "Método de Pago", "Metodo de Pago"],
}


def norm(x):
    return re.sub(r"[^a-z0-9_]", "", str(x).strip().lower().replace(" ", "_"))


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return 0.0


def _excel_col(ref):
    n = 0
    for ch in ref or "":
        if ch.isalpha():
            n = n * 26 + ord(ch.upper()) - 64
        else:
            break
    return n


_XML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _shared_strings(z):
    name = "xl/sharedStrings.xml"
    if name not in z.namelist():
        return []
    out = []
    for _, e in ET.iterparse(z.open(name), events=("end",)):
        if e.tag == _XML_NS + "si":
            out.append("".join((t.text or "") for t in e.iter(_XML_NS + "t")))
            e.clear()
    return out


def _cell_value(c, shared):
    t = c.get("t")
    if t == "inlineStr":
        return "".join((x.text or "") for x in c.iter(_XML_NS + "t"))
    v = c.find(_XML_NS + "v")
    if v is None:
        return None
    x = v.text or ""
    if t == "s":
        try:
            return shared[int(x)]
        except Exception:
            return ""
    if t == "b":
        return x == "1"
    if t == "str":
        return x
    try:
        return float(x)
    except Exception:
        return x


def _excel_date_value(v):
    if v in (None, "", "nan", "None"):
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    try:
        f = float(v)
        # Excel 1900 date system (el usado por los archivos de la app).
        if 20000 <= f <= 80000:
            return (datetime(1899, 12, 30) + timedelta(days=f)).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        pass
    return str(v)


def create_db(path):
    db = sqlite3.connect(path)
    db.execute("PRAGMA journal_mode=OFF")
    db.execute("PRAGMA synchronous=OFF")
    db.execute("PRAGMA temp_store=MEMORY")
    db.execute("PRAGMA cache_size=-65536")
    db.execute("""CREATE TABLE operations(
        id INTEGER PRIMARY KEY,
        archivo TEXT, hoja TEXT, comercio TEXT, fecha TEXT, moneda TEXT,
        recaudo REAL, comision REAL, metodo_pago TEXT,
        Deb_Doc TEXT, Deb_Nombre TEXT, psp_tin TEXT,
        SET_referencia TEXT, Fecha_Transferencia TEXT,
        Banco_Transferencia TEXT, neto REAL, es_negativo INTEGER
    )""")
    return db


def process_to_db(path, db, progress=None):
    """Lector streaming rápido de XLSX vía expat; mapea columnas por nombre de encabezado."""
    wanted_names = {
        "comercio": ALIASES["comercio"], "fecha": ALIASES["fecha"],
        "moneda": ["Moneda", "MONEDA", "Currency"], "recaudo": ALIASES["recaudo"],
        "comision": ALIASES["comision"], "metodo_pago": ALIASES["metodo_pago"],
        "Deb_Doc": ["Deb_Doc"], "Deb_Nombre": ["Deb_Nombre"], "psp_tin": ["psp_tin"],
        "SET_referencia": ["SET_referencia"], "Fecha_Transferencia": ["Fecha Transferencia"],
        "Banco_Transferencia": ["Banco Transferencia"],
    }
    total = 0
    with zipfile.ZipFile(path) as z:
        shared = _shared_strings(z)
        sheets = [x for x in z.namelist() if x.startswith("xl/worksheets/sheet") and x.endswith(".xml")]
        for sheet_no, name in enumerate(sorted(sheets), 1):
            batch = []
            row_num = 0
            current = {}
            colmap = {}
            target_letters = set()
            capturing = False

            def start(el, attrs):
                nonlocal row_num, current, capturing
                if el == "row":
                    row_num += 1
                    current = {}
                elif el == "c":
                    current["__cell__"] = [attrs.get("r"), attrs.get("t"), []]
                elif el in ("v", "t") and "__cell__" in current:
                    capturing = True

            def chars(data):
                if capturing and "__cell__" in current:
                    current["__cell__"][2].append(data)

            def finish(el):
                nonlocal capturing, colmap, total, target_letters
                if el in ("v", "t"):
                    capturing = False
                elif el == "c":
                    c = current.pop("__cell__", None)
                    if c:
                        ref, typ, parts = c
                        raw = "".join(parts)
                        letters = ref.rstrip("0123456789")
                        if row_num > 1 and letters not in target_letters:
                            return
                        col = letters if row_num > 1 else _excel_col(ref)
                        if typ == "s":
                            try:
                                value = shared[int(raw)]
                            except Exception:
                                value = ""
                        else:
                            value = raw
                        current[col] = value
                elif el == "row":
                    if row_num == 1:
                        nh = {norm(v): k for k, v in current.items() if v not in (None, "")}
                        for key, aliases in wanted_names.items():
                            colmap[key] = next((nh[norm(a)] for a in aliases if norm(a) in nh), None)
                        target_letters = set()
                        for vv in colmap.values():
                            if vv is not None:
                                target_letters.add(get_column_letter(vv))
                        colmap = {k: (get_column_letter(v) if v is not None else None) for k, v in colmap.items()}
                        if not (colmap.get("comercio") and colmap.get("recaudo") and colmap.get("comision")):
                            raise ValueError(f"Faltan columnas requeridas en {os.path.basename(path)}")
                    else:
                        def gv(k):
                            return current.get(colmap.get(k)) if colmap.get(k) else None
                        comercio = str(gv("comercio") or "").strip()
                        moneda = str(gv("moneda") or "").strip().upper()
                        rawrec = gv("recaudo")
                        if not comercio and rawrec in (None, ""):
                            current.clear()
                            return
                        rec = num(rawrec)
                        com = num(gv("comision"))
                        psp = str(gv("psp_tin") or "").strip()
                        batch.append((
                            os.path.basename(path), f"sheet{sheet_no}", comercio,
                            _excel_date_value(gv("fecha")), moneda, rec, com,
                            str(gv("metodo_pago") or "").strip(), str(gv("Deb_Doc") or ""),
                            str(gv("Deb_Nombre") or ""), psp, str(gv("SET_referencia") or ""),
                            _excel_date_value(gv("Fecha_Transferencia")), str(gv("Banco_Transferencia") or ""),
                            rec - com, int(rec < 0 or com < 0),
                        ))
                        total += 1
                        if len(batch) >= 10000:
                            db.executemany(
                                'INSERT INTO operations (archivo,hoja,comercio,fecha,moneda,recaudo,comision,'
                                'metodo_pago,Deb_Doc,Deb_Nombre,psp_tin,SET_referencia,Fecha_Transferencia,'
                                'Banco_Transferencia,neto,es_negativo) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                                batch)
                            batch.clear()
                            if progress:
                                progress(total)
                    current.clear()

            parser = expat.ParserCreate()
            parser.StartElementHandler = start
            parser.EndElementHandler = finish
            parser.CharacterDataHandler = chars
            with z.open(name) as stream:
                while True:
                    chunk = stream.read(4 * 1024 * 1024)
                    if not chunk:
                        break
                    parser.Parse(chunk, False)
                parser.Parse(b"", True)
            if batch:
                db.executemany(
                    'INSERT INTO operations (archivo,hoja,comercio,fecha,moneda,recaudo,comision,'
                    'metodo_pago,Deb_Doc,Deb_Nombre,psp_tin,SET_referencia,Fecha_Transferencia,'
                    'Banco_Transferencia,neto,es_negativo) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    batch)
                batch.clear()
            db.commit()
    return total


def style(wb):
    fill = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        if ws.max_row:
            for c in ws[1]:
                c.fill = fill
                c.font = Font(color="FFFFFF", bold=True)
                c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = {str(c.value): c.column for c in ws[1]} if ws.max_row else {}
        for n in ["FECHA", "Fecha Transferencia"]:
            if n in headers:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, headers[n]).number_format = "dd/mm/yyyy hh:mm:ss"
        for n in ["RECAUDO", "COMISION", "NETO", "total_recaudo", "total_comision", "comision_sin_igv", "total_neto"]:
            if n in headers:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, headers[n]).number_format = '#,##0.00'
        # Solo se muestrean las primeras 100 filas para calcular anchos.
        for col in range(1, ws.max_column + 1):
            vals = [ws.cell(r, col).value for r in range(1, min(ws.max_row, 101) + 1)]
            m = max([len(str(v)) for v in vals if v is not None] or [10])
            ws.column_dimensions[get_column_letter(col)].width = min(max(m + 2, 12), 38)
        # Se evita el overhead de Table de openpyxl en hojas muy grandes.
        if 2 <= ws.max_row <= 100000 and ws.max_column >= 1:
            ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            tn = f"T{list(wb.worksheets).index(ws) + 1}_Table"
            t = Table(displayName=tn, ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)


class Filtro:
    """Construye las cláusulas WHERE — idéntico a los métodos de la clase App original."""

    def __init__(self, currency="Todas"):
        self.currency = currency

    def currency_sql(self):
        """Filtro de moneda tomado directamente de la columna Moneda."""
        c = (self.currency or "Todas").strip().upper()
        return f"UPPER(TRIM(COALESCE(moneda,'')))={c!r}" if c in ("PEN", "USD") else "1=1"

    def test_exclusion_sql(self):
        """Excluye solo las reglas de prueba agregadas manualmente por la usuaria."""
        return ("NOT EXISTS (SELECT 1 FROM test_rules tr WHERE tr.comercio=operations.comercio "
                "AND substr(operations.fecha,1,10) BETWEEN tr.fecha_desde AND tr.fecha_hasta)")

    def calc_where(self):
        """Cálculo neto: una sola positiva por PSP_TIN+Moneda; los negativos/reversos sí descuentan."""
        return f"""{self.currency_sql()} AND {self.test_exclusion_sql()} AND (
            es_negativo=1 OR TRIM(COALESCE(psp_tin,''))='' OR id IN (SELECT id FROM main_ids)
        )"""

    def main_where(self):
        """Detalle principal: excluye negativos y evita duplicar una positiva por PSP_TIN+Moneda."""
        return f"""{self.currency_sql()} AND {self.test_exclusion_sql()} AND es_negativo=0 AND (
            TRIM(COALESCE(psp_tin,''))='' OR id IN (SELECT id FROM main_ids)
        )"""


def build_conn(db_path):
    """Recrea exactamente la misma vista/tablas temporales que el método conn() original."""
    con = sqlite3.connect(db_path, check_same_thread=False)
    con.row_factory = sqlite3.Row
    rows = con.execute("SELECT idx,path FROM sources ORDER BY idx").fetchall()
    selects = []
    for idx, path in rows:
        alias = f"s{idx}"
        con.execute(f"ATTACH DATABASE ? AS {alias}", (path,))
        offset = idx * 1000000000
        selects.append(
            f"SELECT id+{offset} AS id, archivo,hoja,comercio,fecha,moneda,recaudo,comision,metodo_pago,"
            f"Deb_Doc,Deb_Nombre,psp_tin,SET_referencia,Fecha_Transferencia,Banco_Transferencia,neto,es_negativo "
            f"FROM {alias}.operations")
    con.execute("CREATE TEMP VIEW operations AS " + " UNION ALL ".join(selects))
    # Se materializan una sola vez las filas usadas por el cálculo principal.
    con.execute("""CREATE TEMP TABLE main_ids AS
        SELECT CASE
                 WHEN MAX(CASE WHEN comision>0 THEN 1 ELSE 0 END)=1
                   THEN MIN(CASE WHEN comision>0 THEN id END)
                 ELSE MIN(id)
               END AS id
        FROM operations
        WHERE es_negativo=0 AND TRIM(COALESCE(psp_tin,''))<>''
        GROUP BY psp_tin, moneda""")
    con.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_main_ids ON main_ids(id)")
    con.execute("""CREATE TEMP TABLE dup_keys AS
        SELECT psp_tin, moneda, COUNT(*) AS apariciones
        FROM operations
        WHERE TRIM(COALESCE(psp_tin,''))<>''
        GROUP BY psp_tin, moneda
        HAVING COUNT(*)>1""")
    con.execute("CREATE INDEX IF NOT EXISTS idx_dup_keys ON dup_keys(psp_tin,moneda)")
    con.execute("CREATE TABLE IF NOT EXISTS test_rules (id INTEGER PRIMARY KEY AUTOINCREMENT, "
                "comercio TEXT NOT NULL, fecha_desde TEXT NOT NULL, fecha_hasta TEXT NOT NULL)")
    con.commit()
    return con


def query_duplicates(con, filtro, limit=None):
    sql = f"""SELECT o.psp_tin,o.moneda Moneda,c.apariciones,o.comercio,o.Deb_Doc,o.Deb_Nombre,o.fecha FECHA,
    o.recaudo RECAUDO,o.comision COMISION,o.neto NETO
    FROM operations o JOIN dup_keys c ON c.psp_tin=o.psp_tin AND COALESCE(c.moneda,'')=COALESCE(o.moneda,'')
    WHERE {filtro.currency_sql()} ORDER BY o.psp_tin,o.fecha,o.id"""
    if limit:
        sql += f" LIMIT {int(limit)}"
    return pd.read_sql_query(sql, con)


def query_details(con, where, limit=None):
    sql = (f"SELECT fecha FECHA,comercio Com_Nombre,Deb_Doc,Deb_Nombre,psp_tin,metodo_pago 'Método de Pago',"
           f"SET_referencia,Fecha_Transferencia 'Fecha Transferencia',Banco_Transferencia 'Banco Transferencia',"
           f"recaudo RECAUDO,comision COMISION,neto NETO FROM operations WHERE {where} ORDER BY id")
    if limit:
        sql += f" LIMIT {int(limit)}"
    return pd.read_sql_query(sql, con)
